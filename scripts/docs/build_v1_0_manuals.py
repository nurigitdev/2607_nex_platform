#!/usr/bin/env python3
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import tempfile
from typing import Iterable
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
RELEASE_ROOT = ROOT / "docs" / "releases" / "v1.0"
SOURCE_ROOT = RELEASE_ROOT / "source"
DELIVERABLE_ROOT = RELEASE_ROOT / "deliverables"
SCREENSHOT_MANIFEST = RELEASE_ROOT / "assets" / "screenshots" / "manifest.json"
BUILD_MANIFEST = DELIVERABLE_ROOT / "manifest.json"
SOURCE_TAG = "v1.0"
SOURCE_REVISION = "ab7af0fc3b42225d1d71d20d517889d694a01c03"
SOURCE_DATE = datetime(2026, 10, 9, 0, 0, 0)
ZIP_DATE = (2026, 10, 9, 0, 0, 0)
FONT_NAME = "Noto Sans CJK KR"
MONO_FONT_NAME = "Liberation Mono"

MANUALS = (
    (
        "01_system_installation_guide.md",
        "NeX_Platform_v1.0_시스템구축절차서.docx",
        "시스템 구축(설치) 절차서",
        "설치 담당자 및 플랫폼 관리자",
    ),
    (
        "02_system_startup_shutdown_guide.md",
        "NeX_Platform_v1.0_시스템기동종료절차서.docx",
        "시스템 기동·종료 절차서",
        "시스템 운영자",
    ),
    (
        "03_system_operations_guide.md",
        "NeX_Platform_v1.0_시스템운영절차서.docx",
        "시스템 운영 절차서",
        "플랫폼 운영자 및 보안 담당자",
    ),
    (
        "04_user_manual.md",
        "NeX_Platform_v1.0_사용자매뉴얼.docx",
        "시스템 사용자 매뉴얼",
        "NeX Platform 최종 사용자",
    ),
)

IMAGE_PATTERN = re.compile(r"^!\[(?P<caption>[^]]*)]\((?P<path>[^)]+)\)$")
HEADING_PATTERN = re.compile(r"^(?P<level>#{1,6})\s+(?P<title>.+)$")
ORDERED_PATTERN = re.compile(r"^\d+\.\s+(?P<text>.+)$")
LINK_PATTERN = re.compile(r"\[([^]]+)]\(([^)]+)\)")
INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+]\([^)]+\))")


class ManualBuildError(RuntimeError):
    pass


def build_all(*, release_root: Path = RELEASE_ROOT) -> dict[str, object]:
    source_root = release_root / "source"
    deliverable_root = release_root / "deliverables"
    screenshot_manifest_path = release_root / "assets" / "screenshots" / "manifest.json"
    screenshot_manifest = _load_screenshot_manifest(screenshot_manifest_path)
    deliverable_root.mkdir(parents=True, exist_ok=True)

    outputs: list[dict[str, object]] = []
    for source_name, output_name, document_title, audience in MANUALS:
        source_path = source_root / source_name
        output_path = deliverable_root / output_name
        if not source_path.is_file():
            raise ManualBuildError(f"manual source is missing: {source_path}")
        document = build_document(
            source_path,
            document_title=document_title,
            audience=audience,
        )
        _save_reproducible(document, output_path)
        outputs.append(
            {
                "source": source_path.relative_to(release_root).as_posix(),
                "source_sha256": _sha256(source_path),
                "deliverable": output_path.relative_to(release_root).as_posix(),
                "deliverable_sha256": _sha256(output_path),
                "size_bytes": output_path.stat().st_size,
            }
        )

    manifest = {
        "schema_version": "nex_v1_manual_deliverables.v1",
        "source_tag": SOURCE_TAG,
        "source_revision": SOURCE_REVISION,
        "source_date": SOURCE_DATE.date().isoformat(),
        "document_count": len(outputs),
        "screenshot_manifest_sha256": _sha256(screenshot_manifest_path),
        "screenshot_count": screenshot_manifest["capture_count"],
        "outputs": outputs,
    }
    _write_json(deliverable_root / "manifest.json", manifest)
    return manifest


def build_document(
    source_path: Path,
    *,
    document_title: str,
    audience: str,
) -> Document:
    lines = source_path.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("# "):
        raise ManualBuildError(f"manual must start with one H1: {source_path}")
    document = Document()
    _configure_document(document, document_title=document_title)
    _add_cover(document, document_title=document_title, audience=audience)
    _add_revision_history(document)
    _add_toc(document)
    document.add_page_break()
    _render_markdown(document, lines[1:], source_path=source_path)
    return document


def _configure_document(document: Document, *, document_title: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.2
    for name, size, color in (
        ("Title", 26, "16324F"),
        ("Heading 1", 18, "16324F"),
        ("Heading 2", 14, "006D77"),
        ("Heading 3", 11, "7A4E00"),
    ):
        style = styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = MONO_FONT_NAME
        code.font.size = Pt(8.5)
        code._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT_NAME)
        code.paragraph_format.left_indent = Cm(0.5)
        code.paragraph_format.right_indent = Cm(0.5)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(7)
        _shade_paragraph_style(code, "F1F4F6")
    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = FONT_NAME
        caption.font.size = Pt(8.5)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string("425466")
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    props = document.core_properties
    props.title = f"NeX Platform v1.0 {document_title}"
    props.subject = "NeX Platform v1.0 Single-host Docker Compose"
    props.author = "NeX Platform"
    props.keywords = "NeX Platform,v1.0,Docker Compose,운영,사용자 매뉴얼"
    props.created = SOURCE_DATE
    props.modified = SOURCE_DATE

    header = section.header.paragraphs[0]
    header.text = f"NeX Platform v1.0 | {document_title}"
    header.style = styles["Caption"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("NeX Platform v1.0  |  ")
    _add_field(footer, "PAGE")
    footer.add_run(" / ")
    _add_field(footer, "NUMPAGES")


def _add_cover(document: Document, *, document_title: str, audience: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(4.0)
    mark = paragraph.add_run("NeX PLATFORM")
    mark.bold = True
    mark.font.name = FONT_NAME
    mark.font.size = Pt(18)
    mark.font.color.rgb = RGBColor.from_string("006D77")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(1.0)
    title = paragraph.add_run(document_title)
    title.bold = True
    title.font.name = FONT_NAME
    title.font.size = Pt(27)
    title.font.color.rgb = RGBColor.from_string("16324F")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version = paragraph.add_run("Version 1.0")
    version.bold = True
    version.font.size = Pt(14)
    document.add_paragraph()
    metadata = document.add_table(rows=5, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.style = "Table Grid"
    rows = (
        ("기준 태그", SOURCE_TAG),
        ("기준 commit", SOURCE_REVISION[:8]),
        ("지원 구성", "Single-host Docker Compose"),
        ("대상 독자", audience),
        ("보안 등급", "사내 운영 문서"),
    )
    for row, values in zip(metadata.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()
    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.add_run(
        "이 문서는 v1.0 코드 기준선을 설명합니다. 실제 운영 배포 승인은 별도 release governance를 따릅니다."
    ).italic = True
    document.add_page_break()


def _add_revision_history(document: Document) -> None:
    document.add_heading("문서 개정 이력", level=1)
    table = document.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(
        table.rows[0].cells,
        ("버전", "일자", "기준", "내용"),
        strict=True,
    ):
        cell.text = text
        _shade_cell(cell, "DDE9EC")
    for cell, text in zip(
        table.rows[1].cells,
        ("1.0", SOURCE_DATE.date().isoformat(), SOURCE_REVISION[:8], "최초 발행"),
        strict=True,
    ):
        cell.text = text
    document.add_paragraph()


def _add_toc(document: Document) -> None:
    document.add_heading("목차", level=1)
    paragraph = document.add_paragraph()
    _add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u')
    paragraph.add_run("문서를 열었을 때 목차 업데이트가 필요하면 전체 필드를 갱신하십시오.")


def _render_markdown(document: Document, lines: list[str], *, source_path: Path) -> None:
    index = 0
    code_lines: list[str] | None = None
    figure_number = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if code_lines is not None:
            if line.startswith("```"):
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                code_lines = None
            else:
                code_lines.append(line)
            index += 1
            continue
        if line.startswith("```"):
            code_lines = []
            index += 1
            continue
        if not line:
            index += 1
            continue
        heading = HEADING_PATTERN.match(line)
        if heading:
            level = min(len(heading.group("level")), 3)
            document.add_heading(_plain_inline(heading.group("title")), level=level)
            index += 1
            continue
        image = IMAGE_PATTERN.match(line)
        if image:
            figure_number += 1
            _add_image(
                document,
                source_path.parent / image.group("path"),
                caption=f"그림 {figure_number}. {image.group('caption')}",
            )
            index += 1
            continue
        if _is_table_row(line) and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            table_lines = [line]
            index += 2
            while index < len(lines) and _is_table_row(lines[index]):
                table_lines.append(lines[index].rstrip())
                index += 1
            _add_table(document, table_lines)
            continue
        ordered = ORDERED_PATTERN.match(line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, ordered.group("text"))
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, line[2:])
            index += 1
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.right_indent = Cm(0.4)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(7)
            _shade_paragraph(paragraph, "FFF4D6")
            run = paragraph.add_run("주의  ")
            run.bold = True
            _add_inline(paragraph, line[2:])
            index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].rstrip()
            if (
                not candidate
                or candidate.startswith(("#", "- ", "> ", "```", "!["))
                or ORDERED_PATTERN.match(candidate)
                or _is_table_row(candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        _add_inline(paragraph, " ".join(paragraph_lines))
    if code_lines is not None:
        raise ManualBuildError(f"unclosed code fence: {source_path}")


def _add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT_NAME
            run.font.size = Pt(9)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            link = LINK_PATTERN.fullmatch(token)
            assert link is not None
            paragraph.add_run(f"{link.group(1)} ({link.group(2)})")
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _add_image(document: Document, path: Path, *, caption: str) -> None:
    resolved = path.resolve()
    if not resolved.is_file():
        raise ManualBuildError(f"referenced image is missing: {resolved}")
    image = DocxImage.from_file(str(resolved))
    width_px, height_px = image.px_width, image.px_height
    if width_px <= 0 or height_px <= 0:
        raise ManualBuildError(f"referenced image has invalid dimensions: {resolved}")
    max_width_inches = 6.2
    max_height_inches = 7.4
    scale = min(max_width_inches / width_px, max_height_inches / height_px)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(
        str(resolved),
        width=Inches(width_px * scale),
        height=Inches(height_px * scale),
    )
    caption_paragraph = document.add_paragraph(style="Figure Caption")
    caption_paragraph.add_run(caption)


def _add_table(document: Document, lines: list[str]) -> None:
    rows = [_split_table_row(line) for line in lines]
    width = len(rows[0])
    if width < 2 or any(len(row) != width for row in rows):
        raise ManualBuildError("Markdown table columns are inconsistent")
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        for cell, value in zip(table.rows[row_index].cells, values, strict=True):
            cell.text = ""
            _add_inline(cell.paragraphs[0], value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                _shade_cell(cell, "DDE9EC")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    document.add_paragraph()


def _is_table_row(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|")


def _is_table_separator(line: str) -> bool:
    if not _is_table_row(line):
        return False
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _plain_inline(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    return LINK_PATTERN.sub(lambda match: match.group(1), text)


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field = OxmlElement("w:instrText")
    field.set(qn("xml:space"), "preserve")
    field.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, field, separate, end))


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _shade_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _shade_paragraph_style(style, fill: str) -> None:
    properties = style.element.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _save_reproducible(document: Document, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="nex-v1-docx-") as temp_dir:
        intermediate = Path(temp_dir) / "intermediate.docx"
        document.save(intermediate)
        with ZipFile(intermediate, "r") as source, ZipFile(
            output_path, "w", compression=ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                original = source.getinfo(name)
                info = ZipInfo(filename=name, date_time=ZIP_DATE)
                info.compress_type = ZIP_DEFLATED
                info.external_attr = original.external_attr
                info.create_system = original.create_system
                target.writestr(info, source.read(name))


def _load_screenshot_manifest(path: Path) -> dict[str, object]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ManualBuildError(f"screenshot manifest is unavailable: {path}") from exc
    if (
        not isinstance(payload, dict)
        or payload.get("source_tag") != SOURCE_TAG
        or payload.get("source_revision") != SOURCE_REVISION
        or payload.get("private_data_included") is not False
        or payload.get("capture_count") != 7
    ):
        raise ManualBuildError("screenshot manifest does not match v1.0 manual policy")
    captures = payload.get("captures")
    if not isinstance(captures, list) or len(captures) != 7:
        raise ManualBuildError("screenshot manifest capture set is incomplete")
    for capture in captures:
        if not isinstance(capture, dict):
            raise ManualBuildError("screenshot manifest capture is invalid")
        image = path.parent / str(capture.get("filename") or "")
        if not image.is_file() or capture.get("sha256") != _sha256(image):
            raise ManualBuildError(f"screenshot digest mismatch: {image.name}")
    return payload


def _write_json(path: Path, payload: dict[str, object]) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def summary_line(manifest: dict[str, object]) -> str:
    return (
        "v1_manual_deliverables=pass "
        f"tag={manifest['source_tag']} "
        f"documents={manifest['document_count']} "
        f"screenshots={manifest['screenshot_count']}"
    )


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build NeX Platform v1.0 DOCX manuals.")
    parser.add_argument("--release-root", type=Path, default=RELEASE_ROOT)
    parser.add_argument("--summary", action="store_true")
    args = parser.parse_args(list(argv) if argv is not None else None)
    try:
        manifest = build_all(release_root=args.release_root)
    except (ManualBuildError, OSError, ValueError) as exc:
        print(f"v1_manual_deliverables=fail reason={exc}")
        return 1
    print(summary_line(manifest) if args.summary else json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
