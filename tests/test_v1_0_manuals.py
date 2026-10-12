from __future__ import annotations

import base64
import json
from pathlib import Path
import sys
from zipfile import ZipFile

from docx import Document
import pytest


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "docs"))

import build_v1_0_manuals as builder  # noqa: E402


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _release_fixture(root: Path) -> Path:
    release = root / "v1.0"
    source = release / "source"
    screenshots = release / "assets" / "screenshots"
    source.mkdir(parents=True)
    screenshots.mkdir(parents=True)
    captures = []
    for index in range(1, 8):
        filename = f"{index:02d}_fixture.png"
        path = screenshots / filename
        path.write_bytes(PNG_1X1)
        captures.append(
            {
                "filename": filename,
                "description": f"fixture {index}",
                "sha256": builder._sha256(path),
                "size_bytes": len(PNG_1X1),
            }
        )
    (screenshots / "manifest.json").write_text(
        json.dumps(
            {
                "source_tag": builder.SOURCE_TAG,
                "source_revision": builder.SOURCE_REVISION,
                "private_data_included": False,
                "capture_count": 7,
                "captures": captures,
            }
        ),
        encoding="utf-8",
    )
    body = """# Fixture manual

## 목적

본문에는 **강조**, `code`, [링크](README.md)가 있다.

> 승인 없이 배포하지 않는다.

1. 첫 번째 단계
2. 두 번째 단계

- 항목 A
- 항목 B

| 이름 | 값 |
| --- | --- |
| 상태 | PASS |

```bash
echo safe
```

![테스트 화면](../assets/screenshots/01_fixture.png)
"""
    for source_name, *_ in builder.MANUALS:
        (source / source_name).write_text(body, encoding="utf-8")
    return release


def test_build_all_creates_four_reproducible_docx_files(tmp_path: Path) -> None:
    release = _release_fixture(tmp_path)

    first = builder.build_all(release_root=release)
    first_digests = [item["deliverable_sha256"] for item in first["outputs"]]
    second = builder.build_all(release_root=release)

    assert first["document_count"] == 4
    assert first["screenshot_count"] == 7
    assert [item["deliverable_sha256"] for item in second["outputs"]] == first_digests
    assert json.loads((release / "deliverables/manifest.json").read_text()) == second
    for item in second["outputs"]:
        path = release / item["deliverable"]
        with ZipFile(path) as archive:
            assert "word/document.xml" in archive.namelist()
            xml = archive.read("word/document.xml").decode("utf-8")
            assert "Fixture manual" not in xml
            assert "승인 없이 배포하지 않는다" in xml
            assert "테스트 화면" in xml


def test_build_document_rejects_invalid_markdown(tmp_path: Path) -> None:
    source = tmp_path / "bad.md"
    source.write_text("본문만 있음\n", encoding="utf-8")

    with pytest.raises(builder.ManualBuildError, match="must start"):
        builder.build_document(source, document_title="Bad", audience="test")

    source.write_text("# 제목\n\n```bash\necho open\n", encoding="utf-8")
    with pytest.raises(builder.ManualBuildError, match="unclosed code fence"):
        builder.build_document(source, document_title="Bad", audience="test")


def test_build_all_rejects_a_missing_manual_source(tmp_path: Path) -> None:
    release = _release_fixture(tmp_path)
    missing = release / "source" / builder.MANUALS[0][0]
    missing.unlink()

    with pytest.raises(builder.ManualBuildError, match="manual source is missing"):
        builder.build_all(release_root=release)


@pytest.mark.parametrize("mutation", ("private", "revision", "count", "digest"))
def test_screenshot_manifest_fails_closed(tmp_path: Path, mutation: str) -> None:
    release = _release_fixture(tmp_path)
    path = release / "assets/screenshots/manifest.json"
    payload = json.loads(path.read_text())
    if mutation == "private":
        payload["private_data_included"] = True
    elif mutation == "revision":
        payload["source_revision"] = "wrong"
    elif mutation == "count":
        payload["capture_count"] = 6
    else:
        payload["captures"][0]["sha256"] = "0" * 64
    path.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(builder.ManualBuildError):
        builder._load_screenshot_manifest(path)


def test_table_and_summary_helpers() -> None:
    assert builder._is_table_row("| A | B |")
    assert builder._is_table_separator("| --- | :---: |")
    assert builder._split_table_row("| A | B |") == ["A", "B"]
    assert builder._plain_inline("**A** `B` [C](x)") == "A B C"
    assert "documents=4" in builder.summary_line(
        {"source_tag": "v1.0", "document_count": 4, "screenshot_count": 7}
    )


def test_render_helpers_cover_compact_and_multiline_boundaries(tmp_path: Path) -> None:
    source = tmp_path / "manual.md"
    source.write_text("# fixture\n", encoding="utf-8")
    document = Document()
    builder._configure_document(document, document_title="Fixture")
    builder._configure_document(document, document_title="Fixture")

    builder._render_markdown(document, ["single paragraph"], source_path=source)
    builder._render_markdown(document, ["first line", "second line"], source_path=source)
    inline = document.add_paragraph()
    builder._add_inline(inline, "**bold only**")

    assert document.paragraphs[-3].text == "single paragraph"
    assert document.paragraphs[-2].text == "first line second line"
    assert document.paragraphs[-1].text == "bold only"


def test_image_table_separator_and_shading_fail_closed(tmp_path: Path, monkeypatch) -> None:
    document = Document()
    builder._configure_document(document, document_title="Fixture")
    missing = tmp_path / "missing.png"
    with pytest.raises(builder.ManualBuildError, match="referenced image is missing"):
        builder._add_image(document, missing, caption="missing")

    invalid = tmp_path / "invalid.png"
    invalid.write_bytes(b"not-an-image")

    class InvalidDimensions:
        px_width = 0
        px_height = 1

    monkeypatch.setattr(builder.DocxImage, "from_file", lambda _: InvalidDimensions())
    with pytest.raises(builder.ManualBuildError, match="invalid dimensions"):
        builder._add_image(document, invalid, caption="invalid")

    with pytest.raises(builder.ManualBuildError, match="columns are inconsistent"):
        builder._add_table(document, ["| only-one-column |"])
    assert not builder._is_table_separator("not a table")

    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    builder._shade_cell(cell, "FFFFFF")
    builder._shade_cell(cell, "000000")


@pytest.mark.parametrize("mutation", ("captures_not_list", "capture_not_object"))
def test_screenshot_manifest_rejects_invalid_capture_shapes(
    tmp_path: Path, mutation: str
) -> None:
    release = _release_fixture(tmp_path)
    path = release / "assets/screenshots/manifest.json"
    payload = json.loads(path.read_text())
    if mutation == "captures_not_list":
        payload["captures"] = "invalid"
    else:
        payload["captures"][0] = None
    path.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(builder.ManualBuildError):
        builder._load_screenshot_manifest(path)


def test_main_reports_failure_and_success(tmp_path: Path, capsys) -> None:
    assert builder.main(["--release-root", str(tmp_path), "--summary"]) == 1
    assert "v1_manual_deliverables=fail" in capsys.readouterr().out

    release = _release_fixture(tmp_path)
    assert builder.main(["--release-root", str(release), "--summary"]) == 0
    assert "v1_manual_deliverables=pass" in capsys.readouterr().out


def test_repository_release_package_matches_source_and_deliverable_manifests() -> None:
    release = ROOT / "docs" / "releases" / "v1.0"
    screenshot_manifest = builder._load_screenshot_manifest(
        release / "assets" / "screenshots" / "manifest.json"
    )
    manifest_path = release / "deliverables" / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert manifest["source_tag"] == builder.SOURCE_TAG
    assert manifest["source_revision"] == builder.SOURCE_REVISION
    assert manifest["document_count"] == len(builder.MANUALS) == 4
    assert manifest["screenshot_count"] == screenshot_manifest["capture_count"] == 7
    assert manifest["screenshot_manifest_sha256"] == builder._sha256(
        release / "assets" / "screenshots" / "manifest.json"
    )

    for output in manifest["outputs"]:
        source = release / output["source"]
        deliverable = release / output["deliverable"]
        assert output["source_sha256"] == builder._sha256(source)
        assert output["deliverable_sha256"] == builder._sha256(deliverable)
        assert output["size_bytes"] == deliverable.stat().st_size
        with ZipFile(deliverable) as archive:
            names = set(archive.namelist())
            assert "word/document.xml" in names
            assert "docProps/core.xml" in names
            assert any(name.startswith("word/media/") for name in names)
